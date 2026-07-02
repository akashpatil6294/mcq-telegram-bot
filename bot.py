import os
import tempfile
import json
from dotenv import load_dotenv
import telebot
from telebot import types
import fitz
from groq import Groq
from docx import Document  # For DOCX support
from fpdf import FPDF      # For PDF generation

load_dotenv()

TELEGRAM_TOKEN = os.getenv("TELEGRAM_TOKEN")
GROQ_API_KEY = os.getenv("GROQ_API_KEY")

bot = telebot.TeleBot(TELEGRAM_TOKEN)
client = Groq(api_key=GROQ_API_KEY)

user_data = {}

class PDF(FPDF):
    def header(self):
        self.set_font('Arial', 'B', 15)
        self.cell(0, 10, 'MCQ Quiz', 0, 1, 'C')
        self.ln(10)

def extract_text(file_path, ext):
    if ext == '.pdf':
        doc = fitz.open(file_path)
        text = "\n".join(page.get_text("text") for page in doc)
        doc.close()
    elif ext == '.docx':
        doc = Document(file_path)
        text = "\n".join([para.text for para in doc.paragraphs])
    else:
        with open(file_path, encoding='utf-8') as f:
            text = f.read()
    return text.strip()

def generate_mcqs(text, num_questions=5, difficulty="medium"):
    prompt = f"""Create exactly {num_questions} {difficulty} level MCQs. Return only valid JSON.

{{
  "questions": [
    {{
      "question": "Question text?",
      "options": ["A. Option 1", "B. Option 2", "C. Option 3", "D. Option 4"],
      "correct": "A",
      "explanation": "Short explanation"
    }}
  ]
}}

Text: {text[:13000]}"""

    response = client.chat.completions.create(
        model="llama-3.3-70b-versatile",
        messages=[{"role": "user", "content": prompt}],
        temperature=0.7,
        max_tokens=4000,
        response_format={"type": "json_object"}
    )
    return response.choices[0].message.content

@bot.message_handler(commands=['start', 'help'])
def start_help(message):
    markup = types.InlineKeyboardMarkup(row_width=1)
    markup.add(types.InlineKeyboardButton("Easy", callback_data="diff_easy"))
    markup.add(types.InlineKeyboardButton("Medium", callback_data="diff_medium"))
    markup.add(types.InlineKeyboardButton("Hard", callback_data="diff_hard"))
    
    bot.reply_to(message, "👋 **MCQ Generator Bot**\n\nChoose difficulty level:", reply_markup=markup)

@bot.callback_query_handler(func=lambda call: call.data.startswith("diff_"))
def difficulty_handler(call):
    diff = call.data.split("_")[1]
    user_data[call.message.chat.id] = {"difficulty": diff, "num": 5}
    bot.answer_callback_query(call.id)
    bot.send_message(call.message.chat.id, f"✅ **{diff.capitalize()}** difficulty selected.\n\nNow send a PDF, DOCX or TXT file.")

@bot.message_handler(content_types=['document'])
def handle_doc(message):
    doc = message.document
    if doc.file_size > 15*1024*1024:
        return bot.reply_to(message, "❌ File too large (max 15MB)")

    file_info = bot.get_file(doc.file_id)
    ext = os.path.splitext(doc.file_name.lower())[1]

    user_pref = user_data.get(message.chat.id, {"difficulty": "medium", "num": 5})
    num = user_pref["num"]
    diff = user_pref["difficulty"]

    bot.reply_to(message, f"🔄 Generating {num} {diff} MCQs...")

    with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
        local_path = tmp.name
        file_data = bot.download_file(file_info.file_path)
        with open(local_path, 'wb') as f:
            f.write(file_data)

    text = extract_text(local_path, ext)
    os.unlink(local_path)

    if len(text) < 100:
        return bot.reply_to(message, "❌ Not enough text found.")

    try:
        result = generate_mcqs(text, num, diff)
        data = json.loads(result)

        # Create PDF
        pdf = PDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        pdf.multi_cell(0, 10, f"MCQ Quiz - {diff.capitalize()} Level\n\n")

        for i, q in enumerate(data.get("questions", []), 1):
            pdf.multi_cell(0, 10, f"Q{i}. {q.get('question')}\n")
            for opt in q.get("options", []):
                pdf.multi_cell(0, 10, opt)
            pdf.multi_cell(0, 10, f"Correct: {q.get('correct')}\nExplanation: {q.get('explanation', '')}\n\n")

        pdf_path = f"mcq_quiz_{message.chat.id}.pdf"
        pdf.output(pdf_path)

        # Send PDF
        with open(pdf_path, 'rb') as f:
            bot.send_document(message.chat.id, f, caption=f"✅ Here is your {num} {diff} MCQs")

        os.unlink(pdf_path)

    except Exception as e:
        bot.reply_to(message, f"❌ Error: {str(e)}")

if __name__ == "__main__":
    print("🚀 Advanced MCQ Bot is Running 24/7...")
    bot.infinity_polling()
